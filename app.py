from flask import Flask, render_template, request, send_file, flash, redirect, url_for, jsonify
import pandas as pd
import os
import io
from datetime import datetime
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

app = Flask(__name__)
app.secret_key = os.environ.get("SECRET_KEY", "mktzap-secret-turn-key")

GV_PESSOAS_DEFAULT = [
    "RAFAELLA CRISTINY JACINTO CAMILO",
    "VANESSA DUARTE",
    "MARCIO MIRANDA ALVES DOS SANTOS",
    "WILLIANE BORGES",
    "LUIS GUILHERME TEODORO DOS SANTOS",
    "ADRIANA CONCEICAO DE OLIVEIRA",
    "DANIELA COELHO"
]

def format_number(val):
    try:
        n = int(val)
        if 0 <= n < 10:
            return f"0{n}"
        return str(n)
    except:
        return "00"

def get_auto_shift_period():
    now_hour = datetime.now().hour
    if now_hour >= 18:
        return "08h às 00h"
    else:
        return "08h às 16h"

@app.route("/", methods=["GET", "POST"])
def home():
    dados_setores = None
    gv_antecipada = None
    total_finalizados = 0
    total_chamados = 0
    top_funcoes = []
    demais_count = 0
    demais_casos = 0
    taxa_geral = 0.0
    canal_lider = "-"
    data_atual = datetime.now().strftime("%d/%m/%Y")
    periodo_turno = get_auto_shift_period()
    nome_arquivo = ""
    erro_formato = None

    if request.method == "POST":
        file = request.files.get("arquivo")

        if not file or file.filename == "":
            flash("Nenhum arquivo foi selecionado.", "error")
            return redirect(url_for("home"))

        nome_arquivo = file.filename
        extensao = os.path.splitext(nome_arquivo)[1].lower()

        formatos_validos = [".xls", ".xlsx", ".html", ".htm", ".csv", ".tsv"]
        if extensao not in formatos_validos:
            erro_formato = f"O formato '{extensao}' não é suportado. Por favor, envie uma planilha em .xls, .xlsx, .html ou .csv."
            return render_template(
                "index.html",
                erro_formato=erro_formato,
                data_atual=data_atual,
                periodo_turno=periodo_turno
            )

        try:
            df = None
            if extensao in [".html", ".htm"]:
                tabelas = pd.read_html(file)
                df = tabelas[0]
            elif extensao in [".xlsx"]:
                df = pd.read_excel(file, header=None)
            elif extensao in [".csv", ".tsv"]:
                df = pd.read_csv(file, header=None, sep=None, engine="python")
            else:
                try:
                    tabelas = pd.read_html(file)
                    df = tabelas[0]
                except Exception:
                    file.seek(0)
                    df = pd.read_excel(file, header=None)

            # Localizar linha de cabeçalho
            header_row = 1 if len(df) > 1 else 0
            for i in range(min(5, len(df))):
                linha_str = " ".join(df.iloc[i].astype(str).str.upper())
                if "SETOR FIM" in linha_str or "TEMPO DE OPERA" in linha_str:
                    header_row = i
                    break

            df.columns = df.iloc[header_row]
            df = df.iloc[header_row + 1:].reset_index(drop=True)
            df.columns = df.columns.astype(str).str.strip()

            col_tempo = next((c for c in df.columns if "TEMPO DE OPERA" in c.upper() or "STATUS" in c.upper()), "Tempo de Operação")
            col_resp = next((c for c in df.columns if "RESPONS" in c.upper() or "OPERADOR" in c.upper()), "Responsável")
            col_setor = next((c for c in df.columns if "SETOR FIM" in c.upper() or "SETOR" in c.upper()), "Setor Fim")
            col_funcao = next((c for c in df.columns if "FUNÇÃO" in c.upper() or "FUNCAO" in c.upper() or "MOTIVO" in c.upper()), "Função do Sistema")
            col_chamado = next((c for c in df.columns if "CHAMADO" in c.upper() or "TICKET" in c.upper()), "CHAMADO")

            df[col_tempo] = df[col_tempo].astype(str).str.strip().str.upper()
            df[col_resp] = df[col_resp].astype(str).str.strip().str.upper()
            df[col_setor] = df[col_setor].astype(str).str.strip()
            df[col_funcao] = df[col_funcao].astype(str).str.strip().str.upper()

            df_finalizados = df[~df[col_tempo].str.contains("NÃO FINALIZADO|NAO FINALIZADO|CANCELADO", na=False)]

            setores = ["Whats GN", "Whats GV", "CB Commerce", "Whats Franquia"]
            df_filtrado = df_finalizados[df_finalizados[col_setor].isin(setores)]

            dados_setores = {}
            for setor in setores:
                base = df_filtrado[df_filtrado[col_setor] == setor]
                chamados_count = 0
                if col_chamado in base.columns:
                    chamados_count = int(base[col_chamado].notna().sum() - (base[col_chamado].astype(str).str.strip().isin(["", "0", "-", "NAN", "NULL", "NÃO", "NAO"])).sum())
                    if chamados_count < 0:
                        chamados_count = 0

                dados_setores[setor] = {
                    "total": len(base),
                    "chamados": chamados_count,
                    "pendencias": 0,
                    "percentual": 0.0
                }

            equipe_form = request.form.get("equipe_integrantes", "")
            gv_pessoas_ativas = [p.strip().upper() for p in equipe_form.split(",") if p.strip()] if equipe_form else GV_PESSOAS_DEFAULT

            df_gv = df_finalizados[
                (df_finalizados[col_setor].str.contains("GV Antecipada|Antecipada", case=False, na=False)) &
                (df_finalizados[col_resp].isin(gv_pessoas_ativas))
            ]

            gv_chamados = 0
            if col_chamado in df_gv.columns:
                gv_chamados = int(df_gv[col_chamado].notna().sum() - (df_gv[col_chamado].astype(str).str.strip().isin(["", "0", "-", "NAN", "NULL", "NÃO", "NAO"])).sum())
                if gv_chamados < 0:
                    gv_chamados = 0

            gv_antecipada = {
                "total": len(df_gv),
                "chamados": gv_chamados,
                "pendencias": 0,
                "percentual": 0.0
            }

            total_finalizados = sum(d["total"] for d in dados_setores.values()) + gv_antecipada["total"]
            total_chamados = sum(d["chamados"] for d in dados_setores.values()) + gv_antecipada["chamados"]
            taxa_geral = (total_chamados / total_finalizados * 100) if total_finalizados > 0 else 0.0

            if total_finalizados > 0:
                for s in dados_setores:
                    dados_setores[s]["percentual"] = (dados_setores[s]["total"] / total_finalizados) * 100
                gv_antecipada["percentual"] = (gv_antecipada["total"] / total_finalizados) * 100

            todos_canais = list(dados_setores.items()) + [("Whats GV Antecipada", gv_antecipada)]
            canal_lider = max(todos_canais, key=lambda x: x[1]["total"])[0]

            serie = df_filtrado[col_funcao].value_counts()
            top3 = serie.head(3)
            resto = serie.iloc[3:]

            top_funcoes = []
            pos = 1
            for nome, valor in top3.items():
                pct = (valor / total_finalizados * 100) if total_finalizados > 0 else 0.0
                top_funcoes.append({
                    "pos": pos,
                    "nome": nome,
                    "quantidade": int(valor),
                    "percentual": round(pct, 1)
                })
                pos += 1

            demais_count = len(resto)
            demais_casos = int(resto.sum()) if len(resto) > 0 else 0

        except Exception as e:
            flash(f"Erro ao processar a planilha: {str(e)}", "error")

    return render_template(
        "index.html",
        dados_setores=dados_setores,
        gv_antecipada=gv_antecipada,
        total_finalizados=total_finalizados,
        total_chamados=total_chamados,
        taxa_geral=round(taxa_geral, 1),
        canal_lider=canal_lider,
        top_funcoes=top_funcoes,
        demais_count=demais_count,
        demais_casos=demais_casos,
        data_atual=data_atual,
        periodo_turno=periodo_turno,
        nome_arquivo=nome_arquivo,
        erro_formato=erro_formato
    )

@app.route("/download_word", methods=["POST"])
def download_word():
    gn_fin = request.form.get("gn_fin", "0")
    gn_cha = request.form.get("gn_cha", "0")
    gn_pen = request.form.get("gn_pen", "0")

    gv_fin = request.form.get("gv_fin", "0")
    gv_cha = request.form.get("gv_cha", "0")
    gv_pen = request.form.get("gv_pen", "0")

    gva_fin = request.form.get("gva_fin", "0")
    gva_cha = request.form.get("gva_cha", "0")
    gva_pen = request.form.get("gva_pen", "0")

    fra_fin = request.form.get("fra_fin", "0")
    fra_cha = request.form.get("fra_cha", "0")
    fra_pen = request.form.get("fra_pen", "0")

    cb_fin = request.form.get("cb_fin", "0")
    cb_cha = request.form.get("cb_cha", "0")
    cb_pen = request.form.get("cb_pen", "0")

    top1 = request.form.get("top1", "1º -: 00")
    top2 = request.form.get("top2", "2º -: 00")
    top3 = request.form.get("top3", "3º -: 00")

    periodo = request.form.get("periodo") or get_auto_shift_period()
    data_doc = request.form.get("data", datetime.now().strftime("%d/%m/%Y"))

    doc = Document()

    p_header = doc.add_paragraph("📋 Passagem de Turno MKTZAP")
    p_header.paragraph_format.space_after = Pt(4)

    p_info = doc.add_paragraph(f"🕗 {periodo} | 📅{data_doc}")
    p_info.runs[0].bold = True
    p_info.paragraph_format.space_after = Pt(12)

    doc.add_paragraph("━━━━━━━━━━━━━━")

    p = doc.add_paragraph("📲 Whats GN")
    p.runs[0].bold = True
    doc.add_paragraph(f"✅ Finalizados: {format_number(gn_fin)}")
    doc.add_paragraph(f"⚙️ Chamados: {format_number(gn_cha)}")
    p_last = doc.add_paragraph(f"⚠️ Pendências: {format_number(gn_pen)}")
    p_last.paragraph_format.space_after = Pt(14)

    p = doc.add_paragraph("📲 Whats GV")
    p.runs[0].bold = True
    doc.add_paragraph(f"✅ Finalizados: {format_number(gv_fin)}")
    doc.add_paragraph(f"⚙️ Chamados: {format_number(gv_cha)}")
    p_last = doc.add_paragraph(f"⚠️ Pendências: {format_number(gv_pen)}")
    p_last.paragraph_format.space_after = Pt(14)

    p = doc.add_paragraph("📲 Whats GV Antecipada")
    p.runs[0].bold = True
    doc.add_paragraph(f"✅ Finalizados: {format_number(gva_fin)}")
    doc.add_paragraph(f"⚙️ Chamados: {format_number(gva_cha)}")
    p_last = doc.add_paragraph(f"⚠️ Pendências: {format_number(gva_pen)}")
    p_last.paragraph_format.space_after = Pt(14)

    p = doc.add_paragraph("📲 Whats Franquia")
    p.runs[0].bold = True
    doc.add_paragraph(f"✅ Finalizados: {format_number(fra_fin)}")
    doc.add_paragraph(f"⚙️ Chamados: {format_number(fra_cha)}")
    p_last = doc.add_paragraph(f"⚠️ Pendências: {format_number(fra_pen)}")
    p_last.paragraph_format.space_after = Pt(14)

    p = doc.add_paragraph("📲 Whats CB Commerce")
    p.runs[0].bold = True
    doc.add_paragraph(f"✅ Finalizados: {format_number(cb_fin)}")
    doc.add_paragraph(f"⚙️ Chamados: {format_number(cb_cha)}")
    p_last = doc.add_paragraph(f"⚠️ Pendências: {format_number(cb_pen)}")
    p_last.paragraph_format.space_after = Pt(18)

    p_top = doc.add_paragraph("🏆 Top Acionamentos\n")
    p_top.runs[0].bold = True

    doc.add_paragraph(top1)
    doc.add_paragraph(top2)
    doc.add_paragraph(top3)

    file_stream = io.BytesIO()
    doc.save(file_stream)
    file_stream.seek(0)

    clean_data = data_doc.replace("/", "-")
    return send_file(
        file_stream,
        as_attachment=True,
        download_name=f"Passagem_de_Turno_MKTZAP_{clean_data}.docx",
        mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )

if __name__ == "__main__":
    port = int(os.environ.get("PORT", 5000))
    app.run(host="0.0.0.0", port=port, debug=False)
